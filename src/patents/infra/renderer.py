from __future__ import annotations

from io import BytesIO
from textwrap import wrap

from patents.adapters import ports


class SimpleTextRenderer(ports.DocumentRenderer):
    file_extension = "txt"

    def render_spec(self, package: ports.FilingPackage) -> bytes:
        lines = [
            f"Title: {package.invention.title}",
            "",
        ]
        for section in package.sections:
            lines.append(f"[{section.section_type.value.upper()}]")
            lines.append(section.text)
            lines.append("")
        return "\n".join(lines).encode("utf-8")

    def render_claims(self, package: ports.FilingPackage) -> bytes:
        lines = [f"{idx + 1}. {claim.text}" for idx, claim in enumerate(package.claims)]
        return "\n".join(lines).encode("utf-8")

    def render_abstract(self, package: ports.FilingPackage) -> bytes:
        abstract = next(
            (s.text for s in package.sections if s.section_type.value == "abstract"),
            package.invention.summary,
        )
        return abstract.encode("utf-8")

    def render_ids(self, package: ports.FilingPackage) -> bytes:
        lines = []
        for art in package.prior_art:
            lines.append(f"{art.title} ({art.publication_date or 'n/a'})")
            if art.url:
                lines.append(art.url)
        return "\n".join(lines).encode("utf-8")


class DocxRenderer(ports.DocumentRenderer):
    file_extension = "docx"

    def render_spec(self, package: ports.FilingPackage) -> bytes:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        _configure_docx(doc, Inches, Pt)
        doc.add_heading(package.invention.title, level=0)
        for section in package.sections:
            heading = section.section_type.value.replace("_", " ").title()
            doc.add_heading(heading, level=1)
            doc.add_paragraph(section.text)
        return _docx_bytes(doc)

    def render_claims(self, package: ports.FilingPackage) -> bytes:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        _configure_docx(doc, Inches, Pt)
        doc.add_heading("Claims", level=0)
        for claim in package.claims:
            doc.add_paragraph(claim.text, style="List Number")
        return _docx_bytes(doc)

    def render_abstract(self, package: ports.FilingPackage) -> bytes:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        _configure_docx(doc, Inches, Pt)
        doc.add_heading("Abstract", level=0)
        abstract = next(
            (s.text for s in package.sections if s.section_type.value == "abstract"),
            package.invention.summary,
        )
        doc.add_paragraph(abstract)
        return _docx_bytes(doc)

    def render_ids(self, package: ports.FilingPackage) -> bytes:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        _configure_docx(doc, Inches, Pt)
        doc.add_heading("Information Disclosure Statement", level=0)
        for art in package.prior_art:
            text = f"{art.title} ({art.publication_date or 'n/a'})"
            if art.url:
                text = f"{text} - {art.url}"
            doc.add_paragraph(text, style="List Bullet")
        return _docx_bytes(doc)


class PdfRenderer(ports.DocumentRenderer):
    file_extension = "pdf"

    def render_spec(self, package: ports.FilingPackage) -> bytes:
        lines = [f"Title: {package.invention.title}", ""]
        for section in package.sections:
            lines.append(section.section_type.value.replace("_", " ").title())
            lines.append(section.text)
            lines.append("")
        return _pdf_bytes(lines)

    def render_claims(self, package: ports.FilingPackage) -> bytes:
        lines = ["Claims", ""]
        lines.extend([f"{idx + 1}. {claim.text}" for idx, claim in enumerate(package.claims)])
        return _pdf_bytes(lines)

    def render_abstract(self, package: ports.FilingPackage) -> bytes:
        abstract = next(
            (s.text for s in package.sections if s.section_type.value == "abstract"),
            package.invention.summary,
        )
        return _pdf_bytes(["Abstract", "", abstract])

    def render_ids(self, package: ports.FilingPackage) -> bytes:
        lines = ["Information Disclosure Statement", ""]
        for art in package.prior_art:
            entry = f"{art.title} ({art.publication_date or 'n/a'})"
            if art.url:
                entry = f"{entry} - {art.url}"
            lines.append(entry)
        return _pdf_bytes(lines)


def _docx_bytes(doc) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _configure_docx(doc, inches_cls, pt_cls) -> None:
    section = doc.sections[0]
    section.top_margin = inches_cls(1)
    section.bottom_margin = inches_cls(1)
    section.left_margin = inches_cls(1)
    section.right_margin = inches_cls(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = pt_cls(12)


def _pdf_bytes(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    width, height = LETTER
    margin = 40
    y = height - margin

    text = pdf.beginText(margin, y)
    text.setFont("Times-Roman", 12)
    for line in lines:
        if line == "":
            text.textLine("")
            continue
        for wrapped in wrap(line, width=90):
            text.textLine(wrapped)
            if text.getY() <= margin:
                pdf.drawText(text)
                pdf.showPage()
                text = pdf.beginText(margin, height - margin)
                text.setFont("Times-Roman", 12)
    pdf.drawText(text)
    pdf.save()
    buffer.seek(0)
    return buffer.read()
